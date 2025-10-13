import os
import re
import threading
from math import ceil
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError

# === CONFIGURATION ===
input_folder = r"C:\xampp\htdocs\alquran\downloads"
base_name = "translations"
sheets_per_file = 50  # create new workbook after every 50 sheets

stop_flag = False

def key_listener():
    """Listen for Q to stop safely."""
    global stop_flag
    while True:
        key = input().strip().upper()
        if key == "Q":
            stop_flag = True
            print("\n🛑 Stop requested by user. Will exit after current file...\n")
            break

def clean_text(value):
    """Remove illegal Excel characters."""
    if not isinstance(value, str):
        return value
    value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F]", "", value)
    value = re.sub(r"[\u200B-\u200F\u202A-\u202E\u2060-\u206F\uFEFF]", "", value)
    return value.strip()

def to_number(value):
    """Convert string digits to int, else keep text."""
    try:
        return int(value)
    except:
        return value

def create_new_workbook(index):
    """Initialize a new workbook."""
    file_name = f"{base_name}_{index}.xlsx"
    output_path = os.path.join(input_folder, file_name)
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["FileName", "SheetName", "RowsAdded"])
    wb.save(output_path)
    wb.close()
    return output_path

def add_txt_to_excel(txt_file, wb_path):
    """Append .txt file content as a new sheet and save."""
    sheet_name = os.path.splitext(txt_file)[0].replace(".", "_")
    txt_path = os.path.join(input_folder, txt_file)

    try:
        wb = load_workbook(wb_path)
    except:
        wb = Workbook()
        wb.active.title = "summary"
        wb.save(wb_path)

    if sheet_name in wb.sheetnames:
        print(f"⏩ Skipped: {sheet_name} (already exists in {os.path.basename(wb_path)})")
        wb.close()
        return False

    try:
        ws = wb.create_sheet(sheet_name)
        ws.append(["Sura", "Ayah", sheet_name])

        with open(txt_path, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        added = 0
        for line in lines:
            parts = line.strip().split("|", 2)
            if len(parts) == 3:
                sura, ayah, text = parts
                ws.append([
                    to_number(clean_text(sura)),
                    to_number(clean_text(ayah)),
                    clean_text(text)
                ])
                added += 1

        if "summary" in wb.sheetnames:
            wb["summary"].append([txt_file, sheet_name, added])

        wb.save(wb_path)
        wb.close()
        print(f"✅ Added: {sheet_name} ({added} rows) → {os.path.basename(wb_path)}")
        return True

    except IllegalCharacterError as e:
        print(f"❌ Illegal char in {sheet_name}: {e}")
    except Exception as e:
        print(f"⚠️ Error processing {txt_file}: {e}")

    return False


# === MAIN SCRIPT ===
print("💡 Press 'Q' anytime + Enter to stop gracefully.\n")
threading.Thread(target=key_listener, daemon=True).start()

txt_files = [f for f in os.listdir(input_folder) if f.endswith(".txt")]
txt_files.sort()

if not txt_files:
    print("⚠️ No .txt files found.")
    exit()

total_files = len(txt_files)
num_workbooks = ceil(total_files / sheets_per_file)
print(f"📚 Found {total_files} .txt files → will create up to {num_workbooks} Excel workbooks.\n")

workbook_index = 1
current_wb_path = create_new_workbook(workbook_index)
sheets_in_current = 0

for i, txt_file in enumerate(txt_files, start=1):
    if stop_flag:
        break

    added = add_txt_to_excel(txt_file, current_wb_path)
    if added:
        sheets_in_current += 1

    # Create new workbook after threshold reached
    if sheets_in_current >= sheets_per_file:
        workbook_index += 1
        current_wb_path = create_new_workbook(workbook_index)
        sheets_in_current = 0

print(f"\n🎉 Done. Created {workbook_index} Excel files under:\n{input_folder}")
