import os
import re
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError

input_folder = "C:/xampp/htdocs/alquran/downloads"
output_excel = os.path.join(input_folder, "translators.xlsx")

def clean_text(value):
    if not isinstance(value, str):
        return value
    value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F]", "", value)
    value = re.sub(r"[\u200B-\u200F\u202A-\u202E\u2060-\u206F\uFEFF]", "", value)
    return value.strip()

# load or create workbook
if os.path.exists(output_excel):
    wb = load_workbook(output_excel)
    print(f"✅ Loaded existing workbook: {output_excel}")
else:
    wb = Workbook()
    ws = wb.active
    ws.title = "translators"
    ws.append(["FileName", "SheetName", "RowsAdded"])
    wb.save(output_excel)
    print(f"🆕 Created new workbook: {output_excel}")

# iterate through text files
for file in os.listdir(input_folder):
    if not file.endswith(".txt"):
        continue

    filepath = os.path.join(input_folder, file)
    sheet_name = os.path.splitext(file)[0].replace(".", "_")

    if sheet_name in wb.sheetnames:
        print(f"⏩ Skipped: {sheet_name} (already exists)")
        continue

    try:
        ws = wb.create_sheet(sheet_name)
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        added = 0
        for line in lines:
            parts = line.strip().split("\t")
            ws.append([clean_text(p) for p in parts])
            added += 1

        # log summary
        if "translators" in wb.sheetnames:
            wb["translators"].append([file, sheet_name, added])
        print(f"✅ Added sheet: {sheet_name} ({added} rows)")

        # save after each sheet for safety
        wb.save(output_excel)
        print(f"💾 Saved after {sheet_name}")

    except IllegalCharacterError as e:
        print(f"❌ Illegal character in {sheet_name}: {e}")
    except Exception as e:
        print(f"⚠️ Error processing {file}: {e}")

print(f"🎉 All done! Workbook: {output_excel}")
