import os
import re
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import IllegalCharacterError

input_folder = r"C:\xampp\htdocs\alquran\downloads"
output_excel = os.path.join(input_folder, "translators.xlsx")

def clean_text(value):
    if not isinstance(value, str):
        return value
    # Remove invisible and control characters
    value = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F]", "", value)
    value = re.sub(r"[\u200B-\u200F\u202A-\u202E\u2060-\u206F\uFEFF]", "", value)
    return value.strip()

def add_sheet_safely(txt_file):
    """Open workbook, add new sheet from txt_file, save, close."""
    sheet_name = os.path.splitext(txt_file)[0].replace(".", "_")
    file_path = os.path.join(input_folder, txt_file)

    # Reopen workbook each time
    if os.path.exists(output_excel):
        wb = load_workbook(output_excel)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "translators"
        ws.append(["FileName", "SheetName", "RowsAdded"])

    # Skip if exists
    if sheet_name in wb.sheetnames:
        print(f"⏩ Skipped: {sheet_name} (already exists)")
        wb.close()
        return

    try:
        ws = wb.create_sheet(sheet_name)
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()

        added = 0
        for line in lines:
            parts = line.strip().split("\t")
            ws.append([clean_text(p) for p in parts])
            added += 1

        # Log entry
        if "translators" in wb.sheetnames:
            wb["translators"].append([txt_file, sheet_name, added])

        # Save and close after each sheet
        wb.save(output_excel)
        wb.close()
        print(f"✅ Added: {sheet_name} ({added} rows) → saved & closed")

    except IllegalCharacterError as e:
        print(f"❌ Illegal character in {sheet_name}: {e}")
    except Exception as e:
        print(f"⚠️ Error processing {txt_file}: {e}")

# --- MAIN LOOP ---
for file in os.listdir(input_folder):
    if not file.endswith(".txt"):
        continue
    add_sheet_safely(file)

print(f"🎉 Completed. All sheets added safely to {output_excel}")
